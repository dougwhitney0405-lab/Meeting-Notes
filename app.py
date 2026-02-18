
import os
import io
import re
from datetime import datetime
from typing import List, Optional

from flask import Flask, render_template, request, send_file, flash, redirect, url_for
from werkzeug.utils import secure_filename

from docx import Document

try:
    from azure.core.credentials import AzureKeyCredential
    from azure.ai.documentintelligence import DocumentIntelligenceClient
except Exception:
    DocumentIntelligenceClient = None
    AzureKeyCredential = None


ALLOWED_DOCX = {"docx"}
ALLOWED_IMG = {"png", "jpg", "jpeg", "webp", "tif", "tiff", "bmp", "heic", "heif"}

def allowed(filename: str, allowed_set: set) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in allowed_set

def extract_text_azure(images: List[bytes]) -> str:
    endpoint = os.getenv("AZURE_DOCUMENT_INTELLIGENCE_ENDPOINT", "").strip()
    key = os.getenv("AZURE_DOCUMENT_INTELLIGENCE_KEY", "").strip()
    if not endpoint or not key or DocumentIntelligenceClient is None:
        return ""

    client = DocumentIntelligenceClient(endpoint=endpoint, credential=AzureKeyCredential(key))

    all_lines = []
    for img_bytes in images:
        poller = client.begin_analyze_document(
            model_id="prebuilt-read",
            analyze_request=img_bytes,
            content_type="application/octet-stream",
        )
        result = poller.result()
        # Collect lines with some structure
        for page in (result.pages or []):
            for line in (page.lines or []):
                if line and line.content:
                    all_lines.append(line.content.strip())
    # de-dup consecutive duplicates
    cleaned = []
    for ln in all_lines:
        if not ln:
            continue
        if cleaned and cleaned[-1] == ln:
            continue
        cleaned.append(ln)
    return "\n".join(cleaned).strip()

def insert_notes_into_docx(doc_bytes: bytes, notes_text: str) -> bytes:
    doc = Document(io.BytesIO(doc_bytes))

    today = datetime.now().strftime("%Y-%m-%d")
    header = f"Handwritten Notes ({today})"

    # Simple bullet formatting from OCR lines
    lines = [ln.strip() for ln in notes_text.splitlines() if ln.strip()]
    # fallback: if OCR returns one big paragraph
    if len(lines) == 1 and len(lines[0]) > 200:
        # split on sentence-ish boundaries
        lines = re.split(r"(?<=[.!?])\s+", lines[0])
        lines = [x.strip() for x in lines if x.strip()]

    # Strategy:
    # 1) If placeholder exists, replace it with a header + bullets.
    placeholder = "[[HANDWRITTEN_NOTES]]"
    replaced = False

    for p in doc.paragraphs:
        if placeholder in p.text:
            p.text = p.text.replace(placeholder, header)
            replaced = True
            # insert bullets after this paragraph
            insert_after = p
            for ln in lines[:500]:
                newp = insert_after.insert_paragraph_after(ln, style="List Bullet")
                insert_after = newp
            break

    if not replaced:
        # Insert after a likely "Agenda" line, otherwise near top after first 2 non-empty paragraphs.
        insert_idx = None
        for i, p in enumerate(doc.paragraphs):
            t = (p.text or "").strip().lower()
            if "agenda" in t or "progress/planning" in t:
                insert_idx = i + 1
                break

        if insert_idx is None:
            # after first 2 non-empty paragraphs
            non_empty = [i for i, p in enumerate(doc.paragraphs) if (p.text or "").strip()]
            insert_idx = (non_empty[1] + 1) if len(non_empty) >= 2 else len(doc.paragraphs)

        # Add at end if index beyond
        if insert_idx >= len(doc.paragraphs):
            doc.add_paragraph(header)
            for ln in lines[:500]:
                doc.add_paragraph(ln, style="List Bullet")
        else:
            # Insert by adding new paragraphs after the paragraph at insert_idx-1
            anchor = doc.paragraphs[max(insert_idx - 1, 0)]
            anchor = anchor.insert_paragraph_after(header)
            for ln in lines[:500]:
                anchor = anchor.insert_paragraph_after(ln, style="List Bullet")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()

app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET_KEY", "dev-secret")

@app.get("/")
def home():
    return render_template("index.html")

@app.post("/process")
def process():
    if "docx" not in request.files:
        flash("Please choose a Word (.docx) file.")
        return redirect(url_for("home"))
    docx_file = request.files["docx"]
    if docx_file.filename == "" or not allowed(docx_file.filename, ALLOWED_DOCX):
        flash("Word file must be .docx")
        return redirect(url_for("home"))

    image_files = request.files.getlist("images")
    image_bytes = []
    for f in image_files:
        if not f or f.filename == "":
            continue
        if not allowed(f.filename, ALLOWED_IMG):
            flash(f"Unsupported image type: {f.filename}")
            return redirect(url_for("home"))
        image_bytes.append(f.read())

    if not image_bytes:
        flash("Please add at least one photo of your handwritten notes.")
        return redirect(url_for("home"))

    doc_bytes = docx_file.read()

    ocr_text = extract_text_azure(image_bytes)
    if not ocr_text:
        # Still produce doc, but include a clear note
        ocr_text = (
            "OCR was not configured or did not return text.\n"
            "To enable handwriting OCR, set AZURE_DOCUMENT_INTELLIGENCE_ENDPOINT and "
            "AZURE_DOCUMENT_INTELLIGENCE_KEY on your hosting service."
        )

    updated = insert_notes_into_docx(doc_bytes, ocr_text)

    out_name = f"Meeting Notes - {datetime.now().strftime('%Y-%m-%d')}.docx"
    return send_file(
        io.BytesIO(updated),
        as_attachment=True,
        download_name=out_name,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

@app.get("/health")
def health():
    return {"ok": True}
